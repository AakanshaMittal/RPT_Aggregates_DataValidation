import os
import re
import json
from pathlib import Path
from collections import defaultdict
from openpyxl import load_workbook
import pandas as pd
import win32com.client
import pythoncom


# =====================================================
# CONFIGURATION
# =====================================================

APP_FOLDER = r"data/inputForModel/File1"

EXCEL_FOLDER = r"data/excel_files"

OUTPUT_FOLDER = r"data/inputForModel/AggregateReport"

MAPPING_FILE = r"config/mapping.json"

SEQUENCE_MAPPING_FILE = r"config/input_for_model_sequence_mapping.json"

TARGET_SHEET = "Input for Model"

TOLERANCE = 0.0001


# =====================================================
# METRIC DEFINITIONS
# =====================================================

METRIC_GROUPS = [
    "Brand Scripts w Exclusions",
    "Total Rebate Collected Per Brand Script",
    "Rebate Guarantees Per Brand Script",
    "Rebate Under/Over Performance"
]

CHANNELS = [
    "Retail 30",
    "Retail 90",
    "Mail",
    "Specialty",
    "Total"
]

YEARS = [
    "Yr. 1",
    "Yr. 2",
    "Yr. 3",
    "Yr. 4",
    "Yr. 5"
]


# =====================================================
# COLUMN LAYOUT
# =====================================================

BLOCK_COLUMNS = {
    "Brand Scripts w Exclusions": (7, 11),       # H:L
    "Total Rebate Collected Per Brand Script": (13, 17),  # N:R
    "Rebate Guarantees Per Brand Script": (19, 23),       # T:X
    "Rebate Under/Over Performance": (25, 29),           # Z:AD
}


# =====================================================
# LOAD MAPPING
# =====================================================

def load_mapping():

    with open(MAPPING_FILE) as f:
        return json.load(f)


def load_sequence_mapping():

    if not os.path.exists(SEQUENCE_MAPPING_FILE):
        return {}

    with open(SEQUENCE_MAPPING_FILE) as f:
        return json.load(f)


def reverse_mapping(mapping):
    return {v: k for k, v in mapping.items()}


# =====================================================
# FILENAME HELPERS
# =====================================================

def extract_agg_id(filename):

    match = re.search(
        r"(AGG\d+[A-Z]{2})",
        filename.upper()
    )

    return match.group(1) if match else None


def extract_tc(filename):

    match = re.search(
        r"TC(\d+[A-Z]?)",
        filename.upper()
    )

    return match.group(1) if match else None


# =====================================================
# NORMALIZATION
# =====================================================

def normalize_value(value):

    if value is None:
        return None

    if pd.isna(value):
        return None

    if isinstance(value, str):

        value = value.strip()

        if value == "":
            return None

        value = value.replace(",", "")

        if value.endswith("%"):

            try:
                return float(value.replace("%", ""))

            except Exception:
                return value

    try:
        return float(value)

    except Exception:
        return str(value).strip()


def values_match(v1, v2):

    n1 = normalize_value(v1)
    n2 = normalize_value(v2)

    if n1 is None and n2 is None:
        return True

    if isinstance(n1, float) and isinstance(n2, float):
        return abs(n1 - n2) <= TOLERANCE

    return str(n1) == str(n2)


# =====================================================
# WORKBOOK READERS
# =====================================================

def recalculate_excel_file(file_path):

    print(
        f"Recalculating workbook: "
        f"{os.path.basename(file_path)}"
    )

    pythoncom.CoInitialize()

    excel = None
    wb = None

    try:

        excel = win32com.client.DispatchEx(
            "Excel.Application"
        )

        excel.Visible = False

        excel.DisplayAlerts = False

        wb = excel.Workbooks.Open(
            os.path.abspath(file_path)
        )

        excel.CalculateFull()

        wb.Save()

        wb.Close(True)

        print(
            f"Recalculation completed: "
            f"{os.path.basename(file_path)}"
        )

    finally:

        try:
            if wb:
                wb.Close(False)
        except:
            pass

        try:
            if excel:
                excel.Quit()
        except:
            pass

        pythoncom.CoUninitialize()

def read_app_file(path):

    print(
        f"Reading APP file: "
        f"{os.path.basename(path)}"
    )

    xls = pd.ExcelFile(
    path,
    engine="openpyxl"
    )
    print(xls.sheet_names)

    return pd.read_excel(
        path,
        sheet_name="Input for Model",
        header=None,
        engine="openpyxl"
    )


def read_excel_workbook(path):

    print(
        f"Reading workbook sheet "
        f"'{TARGET_SHEET}' from "
        f"{os.path.basename(path)}"
    )

    return pd.read_excel(
        path,
        sheet_name=TARGET_SHEET,
        header=None,
        engine="pyxlsb"
    )


# =====================================================
# SAFE CELL ACCESS
# =====================================================

def get_cell(df, row, col):

    try:

        if row >= len(df):
            return None

        if col >= df.shape[1]:
            return None

        return df.iat[row, col]

    except Exception:
        return None


# =====================================================
# RESULT BUILDER
# =====================================================

def build_result(
        section,
        sequence,
        metric,
        channel,
        year,
        app_value,
        excel_value
):

    match = values_match(
        app_value,
        excel_value
    )

    diff = ""

    try:

        if (
            app_value is not None
            and excel_value is not None
        ):

            diff = abs(
                float(app_value)
                -
                float(excel_value)
            )

    except Exception:
        pass

    return {
        "Section": section,
        "Sequence": sequence,
        "Metric": metric,
        "Channel": channel,
        "Year": year,
        "APP_Value": app_value,
        "Excel_Value": excel_value,
        "Difference": diff,
        "Status": "PASS" if match else "FAIL"
    }


# =====================================================
# FILE PAIRING
# =====================================================

def build_file_maps():

    mapping = load_mapping()

    agg_to_tc = reverse_mapping(mapping)

    app_map = {}

    excel_map = {}

    app_path = Path(APP_FOLDER)

    excel_path = Path(EXCEL_FOLDER)

    # -------------------------------
    # APP FILE
    # -------------------------------

    for f in app_path.glob("*.xlsx"):

        agg_id = extract_agg_id(
            f.name
        )

        if agg_id and agg_id in agg_to_tc:

            tc = (
                agg_to_tc[agg_id]
                .replace("TC", "")
            )

            app_map[tc] = f

    # -------------------------------
    # EXCEL WORKBOOK
    # -------------------------------

    for f in excel_path.glob("*.xlsb"):

        tc = extract_tc(
            f.name
        )

        if tc:

            excel_map[tc] = f

    return app_map, excel_map

# =====================================================
# PARSED DATA STRUCTURE
# =====================================================

def make_key(
        section,
        sequence,
        metric,
        channel,
        year
):

    return (
        str(section).strip(),
        str(sequence).strip(),
        str(metric).strip(),
        str(channel).strip(),
        str(year).strip()
    )


# =====================================================
# PARSE COMBINED SECTION
# =====================================================

def parse_combined_section(df):

    parsed = {}

    start_row = 11

    # ---------------------------------------
    # Conservatism
    # ---------------------------------------

    for idx, col in enumerate(range(19, 24)):
        parsed[
        make_key(
            "Combined",
            "",
            "Conservatism",
            "",
            YEARS[idx]
        )] = get_cell(df,8,col)
    # ---------------------------------------
    # Main block
    # ---------------------------------------

    data_rows = {
        "Retail 30": 12,
        "Retail 90": 13,
        "Mail": 14,
        "Specialty": 15,
        "Total": 16
    }

    for metric_name, (
        start_col,
        end_col
    ) in BLOCK_COLUMNS.items():

        for channel, row_num in data_rows.items():

            for idx, col in enumerate(
                range(start_col, end_col + 1)
            ):

                year = YEARS[idx]

                value = get_cell(
                    df,
                    row_num,
                    col
                )

                parsed[
                    make_key(
                        "Combined",
                        "",
                        metric_name,
                        channel,
                        year
                    )
                ] = value

    # ---------------------------------------
    # Total rows below channel section
    # ---------------------------------------

    grand_total_row = 17

    for metric_name, (
        start_col,
        end_col
    ) in BLOCK_COLUMNS.items():

        for idx, col in enumerate(
            range(start_col, end_col + 1)
        ):

            year = YEARS[idx]

            value = get_cell(
                df,
                grand_total_row,
                col
            )

            parsed[
                make_key(
                    "Combined",
                    "",
                    metric_name,
                    "Grand Total",
                    year
                )
            ] = value

    return parsed


# =====================================================
# PARSE SINGLE INDIVIDUAL BLOCK
# =====================================================

def keep_only_app_sequences(
        app_data,
        excel_data
):

    app_sequences = set()

    for key in app_data.keys():

        (
            section,
            sequence,
            metric,
            channel,
            year
        ) = key

        if section == "Individual":
            app_sequences.add(sequence)

    filtered = {}

    for key, value in excel_data.items():

        (
            section,
            sequence,
            metric,
            channel,
            year
        ) = key

        if section != "Individual":

            filtered[key] = value
            continue

        if sequence in app_sequences:

            filtered[key] = value

    return filtered

def parse_individual_block(
        df,
        sequence,
        start_row
):

    parsed = {}

    data_rows = {
        "Retail 30": start_row + 2,
        "Retail 90": start_row + 3,
        "Mail": start_row + 4,
        "Specialty": start_row + 5,
        "Total": start_row + 6
    }

    for metric_name, (
        start_col,
        end_col
    ) in BLOCK_COLUMNS.items():

        for channel, row_num in data_rows.items():

            for idx, col in enumerate(
                range(start_col, end_col + 1)
            ):

                year = YEARS[idx]

                value = get_cell(
                    df,
                    row_num,
                    col
                )

                parsed[
                    make_key(
                        "Individual",
                        sequence,
                        metric_name,
                        channel,
                        year
                    )
                ] = value



    grand_total_row = start_row + 7
    for metric_name, (
            start_col,
            end_col
            ) in BLOCK_COLUMNS.items():
            for idx, col in enumerate(
            range(start_col, end_col + 1)):
                year = YEARS[idx]
                value = get_cell(
                df,
                grand_total_row,
                col
            )
                parsed[
                make_key(
                    "Individual",
                    sequence,
                    metric_name,
                    "Grand Total",
                    year
                )
            ] = value

        

    return parsed


# =====================================================
# FIND INDIVIDUAL BLOCKS
# =====================================================

def parse_individual_rpts(df):

    parsed = {}

    row = 21

    while row < len(df):

        first_val = get_cell(
            df,
            row,
            0
        )

        if str(first_val).strip() == \
                "Guarantees for P&L":

            break

        try:

            seq = int(
                normalize_value(first_val)
            )

        except Exception:

            row += 1

            continue

        block = parse_individual_block(
            df,
            str(seq),
            row
        )

        parsed.update(block)

        row += 9

    return parsed


# =====================================================
# REBATE GUARANTEE SECTION
# =====================================================

def parse_guarantees(df):

    parsed = {}

    start_row = None

    for r in range(len(df)):

        value = str(
            get_cell(
                df,
                r,
                19
            )
        ).strip()

        if value == "REBATE GUARANTEES":

            start_row = r + 1

            break

    if start_row is None:

        return parsed

    row = start_row

    while row < len(df):

        label = get_cell(
            df,
            row,
            19
        )

        value = get_cell(
            df,
            row,
            21
        )

        if label is None:

            row += 1

            continue

        label = str(label).strip()

        if label in ("", "END"):

            row += 1

            continue

        parsed[
            make_key(
                "Guarantee",
                "",
                label,
                "",
                ""
            )
        ] = value

        row += 1

    return parsed


# =====================================================
# MASTER PARSER
# =====================================================

def parse_input_for_model(df):

    parsed = {}

    combined = parse_combined_section(
        df
    )

    parsed.update(combined)

    individuals = parse_individual_rpts(
        df
    )

    parsed.update(individuals)

    guarantees = parse_guarantees(
        df
    )

    parsed.update(guarantees)

    return parsed


# =====================================================
# SEQUENCE REMAP
# =====================================================

def apply_sequence_mapping(
        parsed_data,
        tc,
        sequence_mapping
):

    if tc not in sequence_mapping:
        return parsed_data

    tc_map = sequence_mapping[tc]

    remapped = {}

    for key, value in parsed_data.items():

        (
            section,
            sequence,
            metric,
            channel,
            year
        ) = key

        if section != "Individual":

            remapped[key] = value

            continue

        mapped_sequence = tc_map.get(
            sequence,
            sequence
        )

        new_key = (
            section,
            mapped_sequence,
            metric,
            channel,
            year
        )

        remapped[new_key] = value

    return remapped

# =====================================================
# COMPARE PARSED DATA
# =====================================================

def compare_parsed_data(
        app_data,
        excel_data
):

    results = []

    all_keys = sorted(
        set(app_data.keys()) |
        set(excel_data.keys())
    )

    for key in all_keys:

        app_value = app_data.get(key)

        excel_value = excel_data.get(key)

        (
            section,
            sequence,
            metric,
            channel,
            year
        ) = key

        result = build_result(
            section,
            sequence,
            metric,
            channel,
            year,
            app_value,
            excel_value
        )

        results.append(result)

    return results


# =====================================================
# MISSING KEYS
# =====================================================

def find_missing_keys(
        app_data,
        excel_data
):

    app_keys = set(
        app_data.keys()
    )

    excel_keys = set(
        excel_data.keys()
    )

    missing_in_excel = sorted(
        app_keys - excel_keys
    )

    missing_in_app = sorted(
        excel_keys - app_keys
    )

    return (
        missing_in_excel,
        missing_in_app
    )


# =====================================================
# HTML REPORT
# =====================================================

def generate_html_report(
        results,
        missing_in_excel,
        missing_in_app,
        app_file,
        excel_file,
        output_html
):

    pass_count = sum(
        1 for r in results
        if r["Status"] == "PASS"
    )

    fail_count = sum(
        1 for r in results
        if r["Status"] == "FAIL"
    )

    summary = defaultdict(int)
    for row in results:
        if row["Status"] == "FAIL":
            key = (
            row["Section"],
            row["Metric"]
        )
            summary[key] += 1

    html = []


    html.append("""
    <html>
    <head>
    <style>

    body{
        font-family:Calibri;
    }

    table{
        border-collapse:collapse;
        width:100%;
    }

    th,td{
        border:1px solid black;
        padding:6px;
    }

    th{
        background:#D9D9D9;
    }

    .pass{
        background:#CCFFCC;
    }

    .fail{
        background:#FFCCCC;
    }

    </style>
    </head>

    <body>
    """)

    html.append(
        f"<h1>Input For Model Validation</h1>"
    )

    html.append(
        f"<h2>{os.path.basename(app_file)}"
        f" vs "
        f"{os.path.basename(excel_file)}</h2>"
    )

    html.append(
        f"<p><b>Total Checks:</b> "
        f"{len(results)}</p>"
    )

    html.append(
        f"<p><b>PASS:</b> "
        f"{pass_count}</p>"
    )

    html.append(
        f"<p><b>FAIL:</b> "
        f"{fail_count}</p>"
    )

    html.append("""
<h3>Mismatch Summary</h3>

<table>
<tr>
    <th>Section</th>
    <th>Metric</th>
    <th>Mismatch Count</th>
</tr>
""")
    for (
    section,
    metric
), count in sorted(summary.items()):
        html.append(
        f"""
        <tr>
            <td>{section}</td>
            <td>{metric}</td>
            <td>{count}</td>
        </tr>
        """
    )
    html.append("</table>")

    html.append("""
    <h3>Detailed Results</h3>

    <table>

    <tr>
        <th>Section</th>
        <th>App_Sequence</th>
        <th>Metric</th>
        <th>Channel</th>
        <th>Year</th>
        <th>APP Value</th>
        <th>Excel Value</th>
        <th>Status</th>
    </tr>
    """)

    for row in results:

        if row["Status"] != "FAIL": 
            continue

        css = (
            "pass"
            if row["Status"] == "PASS"
            else "fail"
        )

        html.append(
            f"""
            <tr class="{css}">
                <td>{row['Section']}</td>
                <td>{row['Sequence']}</td>
                <td>{row['Metric']}</td>
                <td>{row['Channel']}</td>
                <td>{row['Year']}</td>
                <td>{row['APP_Value']}</td>
                <td>{row['Excel_Value']}</td>
                <td>{row['Status']}</td>
            </tr>
            """
        )

    html.append("</table>")

    html.append("<h3>Missing in Excel</h3>")
    html.append("<ul>")

    for item in missing_in_excel:
        html.append(
            f"<li>{item}</li>"
        )

    html.append("</ul>")

    html.append("<h3>Missing in APP</h3>")
    html.append("<ul>")

    for item in missing_in_app:
        html.append(
            f"<li>{item}</li>"
        )

    html.append("</ul>")

    html.append(
        "</body></html>"
    )

    with open(
            output_html,
            "w",
            encoding="utf-8"
    ) as f:

        f.write(
            "\n".join(html)
        )


# =====================================================
# CONSOLIDATED REPORT
# =====================================================

def generate_consolidated_report(
        summary_rows,
        output_file
):

    df = pd.DataFrame(
        summary_rows
    )

    with pd.ExcelWriter(
        output_file,
        engine="openpyxl"
    ) as writer:

        df.to_excel(
            writer,
            index=False,
            sheet_name="Summary"
        )


# =====================================================
# PROCESS SINGLE TC
# =====================================================

def process_tc(
        tc,
        app_file,
        excel_file,
        sequence_mapping
):

    print(
        f"\nProcessing TC{tc}"
    )

    
    wb = load_workbook(
    app_file,
    data_only=False
    )
    ws = wb["Input for Model"]
    '''print("\n===== FORMULA DEBUG =====")
    for cell in ["N13","O13","P13","T13","U13","V13","Z13","AA13" ]:
        print(cell, ws[cell].value)'''

    wb2 = load_workbook(
    app_file,
    data_only=True)
    ws2 = wb2["Input for Model"]
    '''print("\n===== VALUE CHECK =====")

    for cell in ["N13","H13","T13","Z13"]:
        print(cell, ws2[cell].value)

    print("\n===== CACHED VALUES =====")
    
    for cell in ["N13","H13","T13","T13","U13","V13","Z13","AA13" ]:
        print(cell, ws2[cell].value)'''
    

    recalculate_excel_file(
    app_file)

    app_df = read_app_file(
    app_file
    )
    '''print("\n===== APP DATAFRAME =====")
    print("APP SHAPE:", app_df.shape)
    print(app_df.iloc[0:40, 0:10])

    print("APP SHAPE:", app_df.shape)
    for i in range(min(40, len(app_df))):
        print(i, app_df.iloc[i].tolist())'''

    excel_df = read_excel_workbook(
        excel_file
    )

    '''print("\n===== EXCEL DATAFRAME =====")
    print("EXCEL SHAPE:", excel_df.shape)
    print(excel_df.iloc[0:40, 0:10])

    print("EXCEL SHAPE:", excel_df.shape)
    for i in range(min(40, len(app_df))):
        print(i, app_df.iloc[i].tolist())'''

    app_data = parse_input_for_model(
        app_df
    )

    excel_data = parse_input_for_model(
        excel_df
    )

    excel_data = apply_sequence_mapping(excel_data, f"TC{tc}",sequence_mapping)
    excel_data = keep_only_app_sequences(app_data, excel_data)

    results = compare_parsed_data(
        app_data,
        excel_data
    )

    (
        missing_in_excel,
        missing_in_app
    ) = find_missing_keys(
        app_data,
        excel_data
    )

    os.makedirs(
        OUTPUT_FOLDER,
        exist_ok=True
    )

    output_html = os.path.join(
        OUTPUT_FOLDER,
        f"TC{tc}_InputForModel.html"
    )

    generate_html_report(
        results,
        missing_in_excel,
        missing_in_app,
        str(app_file),
        str(excel_file),
        output_html
    )

    fail_cnt = sum(
        1
        for r in results
        if r["Status"] == "FAIL"
    )

    return {
        "TC": f"TC{tc}",
        "Checks": len(results),
        "Failures": fail_cnt,
        "Status":
            "PASS"
            if fail_cnt == 0
            else "FAIL"
    }


# =====================================================
# MAIN
# =====================================================

def main():

    os.makedirs(
        OUTPUT_FOLDER,
        exist_ok=True
    )

    sequence_mapping = (
        load_sequence_mapping()
    )

    app_map, excel_map = (
        build_file_maps()
    )

    print(
        "\nAPP MAP:"
    )

    print(app_map)

    print(
        "\nEXCEL MAP:"
    )

    print(excel_map)

    common_tcs = sorted(
        set(app_map.keys()) &
        set(excel_map.keys())
    )

    print(
        f"\nRunning "
        f"{len(common_tcs)} "
        f"scenario(s)\n"
    )

    summary_rows = []

    for tc in common_tcs:

        summary = process_tc(
            tc,
            app_map[tc],
            excel_map[tc],
            sequence_mapping
        )

        summary_rows.append(
            summary
        )

    consolidated_file = os.path.join(
        OUTPUT_FOLDER,
        "Consolidated_InputForModel.xlsx"
    )

    generate_consolidated_report(
        summary_rows,
        consolidated_file
    )

    print(
        "\nCompleted."
    )


# =====================================================
# ENTRY
# =====================================================

if __name__ == "__main__":
    main()